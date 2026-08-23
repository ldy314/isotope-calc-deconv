# -*- coding: utf-8 -*-
"""端到端测试：COM(AutomationSecurity=1) 打开 xlsm，触发 RunX（无 MsgBox），
再读回结果工作表（openpyxl data_only）核对。"""
import os
import time
import win32com.client as wc
import openpyxl

os.environ["ISO_NO_MSGBOX"] = "1"   # 让宏跳过模态 MsgBox，写状态到 输入!D3

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, "同位素杂质计算.xlsm")
base = os.path.splitext(os.path.basename(XLSX))[0]
folder = HERE

excel = wc.DispatchEx("Excel.Application")
excel.Visible = False
excel.DisplayAlerts = False
excel.AutomationSecurity = 1
try:
    wb = excel.Workbooks.Open(XLSX, 0, False)
    has_runx = False
    for comp in wb.VBProject.VBComponents:
        if comp.Name == "ModRun":
            has_runx = "Sub RunX" in comp.CodeModule.Lines(1, comp.CodeModule.CountOfLines)
    print("ModRun/RunX present:", has_runx)
    t0 = time.time()
    excel.Run("RunX")
    print("RunX returned in %.1fs" % (time.time() - t0))
    wb.Close(False)
finally:
    excel.Quit()

print("\n--- 读回结果 (data_only) ---")
wb2 = openpyxl.load_workbook(XLSX, data_only=True)
ws_in = wb2["输入"]
print("状态(D3):", ws_in["D3"].value)

for sh in ("四法总览", "z3_档汇总", "z3_杂质明细", "校准与实测"):
    if sh not in wb2.sheetnames:
        print(f"[{sh}] MISSING"); continue
    ws = wb2[sh]
    print(f"\n[{sh}]")
    n = 0
    for row in ws.iter_rows(values_only=True):
        if all(v is None for v in row):
            continue
        print("  ", row)
        n += 1
        if n >= 12:
            break

for fn in os.listdir(folder):
    if fn.startswith(base + "_") and (fn.endswith(".csv") or fn.endswith(".txt")):
        try:
            os.remove(os.path.join(folder, fn))
        except Exception:
            pass
print("\ncleaned temp csv/txt.")
