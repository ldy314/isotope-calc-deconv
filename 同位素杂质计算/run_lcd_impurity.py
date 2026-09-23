# -*- coding: utf-8 -*-
"""
run_lcd_impurity.py — 对岛津 Q-TOF .lcd 计算「同位素杂质含量」（9 档模型）

与 run_all.py 的分工
--------------------
run_all.py 面向已导出的 **mzML**（sp_003 / STD 0.005 那批，报告文案已固化为该批）；
本脚本面向 **.lcd 原始文件**，自带 TOF 质量轴标定，复用同一套 9 档逐级三角扣除引擎
（calc.analyze_profile），可对任意批次跑 STD / SP。

关键点（2026-09-23）
------------------
1. `.lcd`（QTFL Centroid）存储值 x 与 m/z 满足 m/z = a·x² + c，**不是**线性换算；
   a、c 由 6 个已知锚点最小二乘标定（残差 ±10 ppm）。
2. 谱为**质心谱**，分箱后单个真实峰会被打散成多个相邻 bin →
   窗内「积分」会把弱峰相对放大；对本类数据 **M0·峰顶** 与独立方法
   （解卷积/20260914分析 的包络序号法）吻合更好（t1：2.57% vs 2.52%）。
   故输出四法全量，主值取 **M0·峰顶**，并与 M0·积分并列供核对。

用法
----
    python run_lcd_impurity.py [--data-dir DIR] [--out DIR] [--rt-lo S] [--rt-hi S]
"""
from __future__ import annotations

import argparse
import os
import sys

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
PARENT = os.path.dirname(HERE)
for _p in (PARENT, HERE):
    if _p not in sys.path:
        sys.path.insert(0, _p)

import calc as C
import lcd_io
import model as M

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---- 默认数据与参数 ---------------------------------------------------------
DEFAULT_DATA = os.path.join(PARENT, "计算数据", "20260914 ID of SPH20291")
DEFAULT_SAMPLES = [
    ("STD_005", "STD_005.lcd", "天然形对照（应不含标记本体）"),
    ("ISO1_004", "SYSPH20291-Isotopel-260911_004.lcd", "样品（SPH20291-Isotope1）"),
]
RT_LO, RT_HI = 261.0, 266.0      # = 4.35–4.43 min，覆盖 4.4 min 目标峰
Z_LIST = (3, 4)                   # z=3 主定量、z=4 辅助（z=2 信噪比不足，不报）

# 6 个 TOF 标定锚点：(x_measured, z, tier)；tier0 = 全标记本体、tier8 = 天然形。
# x 取自实测 M0 峰位，理论值由 model.MODEL 给出；来源见 解卷积/20260914分析/analyze_20260914.py。
ANCHORS = [(1764.742, 2, 8), (1441.284, 3, 8), (1248.498, 4, 8),
           (1767.244, 2, 0), (1443.324, 3, 0), (1250.266, 4, 0)]

METHODS = [("m0_top", "M0·峰顶"), ("m0_integral", "M0·积分"),
           ("base_top", "基峰·峰顶"), ("base_integral", "基峰·积分")]
PRIMARY = "m0_top"

# ---- 样式 ------------------------------------------------------------------
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
SUB_FILL = PatternFill("solid", fgColor="D9E1F2")
KEY_FILL = PatternFill("solid", fgColor="FFF2CC")
HEAD_FONT = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
PCT = "0.00"


def put(ws, r, c, v, bold=False, fill=None, align=CENTER, num=None):
    cell = ws.cell(row=r, column=c, value=v)
    if bold:
        cell.font = BOLD
    if fill:
        cell.fill = fill
    cell.alignment = align
    if num:
        cell.number_format = num
    cell.border = BORDER
    return cell


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data-dir", default=DEFAULT_DATA)
    ap.add_argument("--out", default=None, help="输出目录（默认写入数据目录）")
    ap.add_argument("--rt-lo", type=float, default=RT_LO)
    ap.add_argument("--rt-hi", type=float, default=RT_HI)
    args = ap.parse_args()
    out_dir = args.out or args.data_dir
    os.makedirs(out_dir, exist_ok=True)

    # ---- 标定 ----
    anch = [(x, M.MODEL[z][t]["m0_mz"]) for x, z, t in ANCHORS]
    a, c, resid = lcd_io.calibrate_tof(anch)
    print(f"[标定] m/z = {a:.8e}·x² + ({c:+.5f})  残差 ppm = {[round(r, 1) for r in resid]}")

    # ---- 计算 ----
    results = {}
    for tag, fn, note in DEFAULT_SAMPLES:
        path = os.path.join(args.data_dir, fn)
        print("=" * 90)
        if not os.path.exists(path):
            print(f"[跳过] 缺文件 {path}")
            continue
        d = lcd_io.read_lcd(path, args.rt_lo, args.rt_hi, (a, c))
        print(f"{tag}  {fn}  RT {d['rt_lo']:.1f}-{d['rt_hi']:.1f}s "
              f"({d['rt_lo']/60:.2f}-{d['rt_hi']/60:.2f} min)  scans={d['n_scans']} "
              f"峰数={len(d['mz'])}")
        results[tag] = {"file": fn, "note": note, "rt": (d["rt_lo"], d["rt_hi"]),
                        "n_scans": d["n_scans"], "per_z": {}}
        for z in Z_LIST:
            r = C.analyze_profile(d["mz"], d["intensity"], z, tag,
                                  rt_lo=d["rt_lo"], rt_hi=d["rt_hi"], n_scans=d["n_scans"])
            results[tag]["per_z"][z] = r
            cc = r["results"][PRIMARY]["content_pct"]
            print(f"  z={z}  {PRIMARY}: " + "  ".join(f"t{t}={cc[t]:.2f}%" for t in range(9)))

    if not results:
        raise SystemExit("没有可计算的数据")

    # ---- Excel ----
    wb = Workbook()

    # Sheet 1 总览
    ws = wb.active
    ws.title = "结果总览"
    widths(ws, [12, 10, 12] + [9] * 9 + [10])
    ws["A1"] = "同位素杂质含量（9 档模型）— 20260914 ID of SPH20291 批次"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = (f"RT 窗口 {args.rt_lo:.1f}–{args.rt_hi:.1f} s "
                f"({args.rt_lo/60:.2f}–{args.rt_hi/60:.2f} min)｜"
                f"TOF 标定 m/z = {a:.6e}·x² {c:+.5f}（残差 ±{max(abs(r) for r in resid):.0f} ppm）｜"
                f"主值 = M0·峰顶")
    ws["A2"].font = Font(size=9, color="555555")
    hdr = ["样品", "电荷 z", "估计量"] + [f"t{t}" for t in range(9)] + ["备注"]
    for j, h in enumerate(hdr, start=1):
        put(ws, 4, j, h, bold=True, fill=HEAD_FILL)
        ws.cell(row=4, column=j).font = HEAD_FONT
    row = 5
    for tag, info in results.items():
        for z in Z_LIST:
            r = info["per_z"].get(z)
            if not r:
                continue
            for key, lab in METHODS:
                cc = r["results"][key]["content_pct"]
                fill = KEY_FILL if key == PRIMARY else None
                put(ws, row, 1, tag, bold=(key == PRIMARY), fill=fill, align=LEFT)
                put(ws, row, 2, z, fill=fill)
                put(ws, row, 3, lab, bold=(key == PRIMARY), fill=fill, align=LEFT)
                for t in range(9):
                    put(ws, row, 4 + t, round(float(cc[t]), 3), fill=fill, num=PCT)
                put(ws, row, 13, info["note"] if key == PRIMARY else "", align=LEFT)
                row += 1
    ws.freeze_panes = "D5"
    ws["A" + str(row + 1)] = ("t0 = 全标记本体（SPH20291-Isotope1）；t1…t8 = 依次少 1…8 个标记原子"
                              "（未完全标记杂质）。STD_005 为天然形对照：t0/t1 = 0 表示不含标记本体，"
                              "其 t4–t8 分布来自天然同位素包络的投影，非杂质。")
    ws["A" + str(row + 1)].font = Font(size=9, color="C00000")

    # Sheet 2/3 各样品明细
    for tag, info in results.items():
        for z in Z_LIST:
            r = info["per_z"].get(z)
            if not r:
                continue
            ws = wb.create_sheet(f"{tag}_z{z}")
            widths(ws, [7, 52, 11, 11, 11, 11, 12, 12, 12])
            ws["A1"] = f"{tag} — 同位素杂质含量（z={z}）｜{info['file']}"
            ws["A1"].font = Font(bold=True, size=12)
            ws["A2"] = (f"RT {info['rt'][0]:.1f}–{info['rt'][1]:.1f}s｜"
                        f"主值 M0·峰顶（本批质心谱下与独立方法吻合）")
            ws["A2"].font = Font(size=9, color="555555")
            head = ["档", "物种", "M0·峰顶 %", "M0·积分 %", "基峰·峰顶 %", "基峰·积分 %",
                    "理论 M0 m/z", "窗内峰顶", "窗内积分"]
            for j, h in enumerate(head, start=1):
                put(ws, 4, j, h, bold=True, fill=HEAD_FILL)
                ws.cell(row=4, column=j).font = HEAD_FONT
            rows = r["results"][PRIMARY]["rows"]
            mi = r["measured"]["m0"]["integral"]
            mt = r["measured"]["m0"]["top"]
            notes = r["results"][PRIMARY]["notes"]
            rr = 5
            for item in rows:
                t = item["tier"]
                c1 = r["results"]["m0_top"]["content_pct"][t]
                c2 = r["results"]["m0_integral"]["content_pct"][t]
                c3 = r["results"]["base_top"]["content_pct"][t]
                c4 = r["results"]["base_integral"]["content_pct"][t]
                fill = KEY_FILL if t == 0 else None
                name = item["name"]
                if isinstance(item.get("m0_mz_real"), list):
                    name = f"{name}｜m/z " + ", ".join(f"{v:.3f}" for v in item["m0_mz_real"][:4])
                put(ws, rr, 1, t, bold=True, fill=fill)
                put(ws, rr, 2, name, fill=fill, align=LEFT)
                put(ws, rr, 3, round(float(c1), 3), bold=True, fill=fill, num=PCT)
                put(ws, rr, 4, round(float(c2), 3), fill=fill, num=PCT)
                put(ws, rr, 5, round(float(c3), 3), fill=fill, num=PCT)
                put(ws, rr, 6, round(float(c4), 3), fill=fill, num=PCT)
                put(ws, rr, 7, round(float(item["m0_mz_real"][0] if isinstance(
                    item["m0_mz_real"], list) else item["m0_mz_real"]), 4))
                put(ws, rr, 8, round(float(mt[t]), 0), num="#,##0")
                put(ws, rr, 9, round(float(mi[t]), 0), num="#,##0")
                if t in notes:
                    ws.cell(row=rr, column=2).value = f"{name}  ⚠ {notes[t]}"
                rr += 1
            tot = sum(r["results"][PRIMARY]["content_pct"].values())
            put(ws, rr, 2, f"合计（归一化）={tot:.2f}%", bold=True, align=LEFT)
            ws.freeze_panes = "C5"

    # Sheet 4 校准与实测
    ws = wb.create_sheet("校准与实测")
    widths(ws, [12, 6, 5, 7, 13, 13, 11, 13, 13, 12])
    head = ["样品", "z", "档", "锚点", "理论 m/z", "实测峰心", "偏差 ppm",
            "窗内峰顶", "窗内积分", "基线"]
    for j, h in enumerate(head, start=1):
        put(ws, 1, j, h, bold=True, fill=HEAD_FILL)
        ws.cell(row=1, column=j).font = HEAD_FONT
    rr = 2
    for tag, info in results.items():
        for z in Z_LIST:
            r = info["per_z"].get(z)
            if not r:
                continue
            for t in range(9):
                for anchor, lab in (("m0", "M0"), ("base", "基峰")):
                    if anchor == "m0":
                        th = M.MODEL[z][t]["m0_mz"]
                        meas_v = r["measured"]["m0"]["top"][t]
                        meas_i = r["measured"]["m0"]["integral"][t]
                    else:
                        th = M.MODEL[z][t]["base_mz"]
                        meas_v = r["measured"]["base"]["top"][t]
                        meas_i = r["measured"]["base"]["integral"][t]
                    ctr = r["centers"][(t, anchor)]
                    put(ws, rr, 1, tag, align=LEFT)
                    put(ws, rr, 2, z)
                    put(ws, rr, 3, t)
                    put(ws, rr, 4, lab)
                    put(ws, rr, 5, round(float(th), 4))
                    put(ws, rr, 6, round(float(ctr), 4))
                    put(ws, rr, 7, round((ctr - th) * 1e6 / th, 1))
                    put(ws, rr, 8, round(float(meas_v), 0), num="#,##0")
                    put(ws, rr, 9, round(float(meas_i), 0), num="#,##0")
                    put(ws, rr, 10, round(float(r["baselines"].get((t, anchor), 0.0)), 1))
                    rr += 1
    ws.freeze_panes = "A2"

    # Sheet 5 方法与参数
    ws = wb.create_sheet("方法与参数")
    widths(ws, [120])
    lines = [
        "一、数据与换算",
        f"  · 数据目录：{args.data_dir}",
        "  · 岛津 Q-TOF .lcd（QTFL Centroid）：存储值 x 与 m/z 满足 m/z = a·x² + c，不是线性换算；",
        f"    a = {a:.8e}，c = {c:+.5f}（6 锚点最小二乘，残差 ±{max(abs(r) for r in resid):.0f} ppm）",
        f"  · RT 窗 {args.rt_lo:.1f}–{args.rt_hi:.1f} s = {args.rt_lo/60:.2f}–{args.rt_hi/60:.2f} min（4.4 min 目标峰）",
        "  · 累加窗内质心谱，按 x 分箱 1e-3",
        "",
        "二、模型（model.py，9 档）",
        "  · 本体 = SPH20291-Isotope1（C128[13C]6H198N26[15N]2O35S2），共 8 个标记原子",
        "  · 档 t = 有 t 个标记原子被天然同位素替换（t=0 全标记本体，t=8 全天然）",
        "  · 同档内 13C/15N 组合在 5 ppm 下质量简并 → 每档合并为 1 行、含量取整档总量",
        "",
        "三、含量算法（calc.py，逐级三角扣除）",
        "  · 9 档沿质量轴构成二对角链：位置 M0(t) 上只叠『本档 M0』与『更轻一档 t+1 的基峰』",
        "  · 理论重叠系数 W[i][j] = tier j 全部理论峰落在 tier i 锚点 ±0.04 Da 内的概率之和（5 ppm）",
        "  · M0 法：t=8→0 逐级扣除；基峰法：t=0→8 逐级扣除；负值截断为 0",
        "  · 归一化：含量% = 该档量 / Σ(全部 9 档) × 100",
        "",
        "四、为什么主值取 M0·峰顶 而非 M0·积分",
        "  · 本批为**质心谱**，按 x 分箱后单个真实峰被打散成多个相邻 bin",
        "    （实测 tier0 跨 5 个 bin、tier1 跨 10 个 bin），窗内『积分』会相对放大弱峰与噪声；",
        "  · M0·峰顶（窗内最大强度）与独立方法（解卷积/20260914分析 的包络序号法）吻合：",
        "    ISO1_004 轻一档 z=3 2.57%（报告 2.52%）、z=4 1.27%（报告 1.26%）；",
        "  · 四法结果均列出，供交叉核对。",
        "",
        "五、结果解读",
        "  · ISO1_004（SPH20291-Isotope1）：t0 = 全标记本体，t1 = 少 1 个标记原子的杂质（主杂质）。",
        "  · STD_005（天然形对照）：t0/t1 = 0 → 不含标记本体；其 t4–t8 分布来自**天然同位素包络**",
        "    在相应档位的投影，不是真实杂质（该结论与 20260914分析 报告一致）。",
        "  · z=2 因窗内噪声占比高（>70%）不予报出。",
        "",
        "六、与既有交付物的关系",
        "  · 引擎与参数（5 ppm 系数、±0.04 Da 提取窗、四估计量）与 run_all.py 完全一致；",
        "  · 差异仅在于输入端：本脚本直接读 .lcd 并做 TOF 标定，无需预先导出 mzML。",
    ]
    for i, s in enumerate(lines, start=1):
        cell = ws.cell(row=i, column=1, value=s)
        cell.alignment = LEFT
        if s and (s[0] in "一二三四五六"):
            cell.font = BOLD

    out_path = os.path.join(out_dir, "同位素杂质含量_20260914批次.xlsx")
    wb.save(out_path)
    print("=" * 90)
    print(f"[SAVED] {out_path}")
    return out_path


if __name__ == "__main__":
    main()
