# -*- coding: utf-8 -*-
"""
deconv_impurity.py - 基于解卷积的同位素杂质含量计算
输入：岛津 .lcd / .csv 文件 + 标记化合物元素表
输出：xlsx 报表（与 run_all.py 相同格式，1 位小数）

复用已有脚本：
  lcd2csv.py : .lcd → 谱图 CSV
  deconv.py  : NNLS 解卷积（模式拟合）
  imp.py     : 杂质枚举
  theo.py    : 理论同位素包络

支持两种样品类型：
  sp  : SPH20291-Isotope1（全标记本体）C128[13C]6 H198 N26[15N]2 O35 S2
  std : SPH20291（天然本体）C134 H198 N28 O35 S2
"""

from __future__ import annotations
import os
import sys
import argparse
from typing import List, Tuple, Dict, Optional

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))  # 上级目录（含 deconv.py / lcd2csv.py / imp.py / theo.py）

import deconv
import lcd2csv
import imp as imp_mod


# ---------------------------------------------------------------------------
# 常量
# ---------------------------------------------------------------------------
RESOLUTION = 30000                    # LCMS-9030 标称分辨率
PROFILE_BIN_DA = 0.01                 # centroid → synthetic profile 的网格间距
PROFILE_K_SIGMA = 3.0                 # synthetic profile 高斯延伸（每侧 σ 数）
GAUSS_K = 2.354820045                 # FWHM → σ

# SPH20291-Isotope1（全标记本体，样品 sp）
LABELED_COLS = [
    ("C", "natural", 128), ("C", "13", 6),
    ("H", "natural", 198),
    ("N", "natural", 26), ("N", "15", 2),
    ("O", "natural", 35), ("S", "natural", 2),
]
BODY_FORMULA_LABELED = "C₁₂₈¹³C₆H₁₉₈N₂₆¹⁵N₂O₃₅S₂"

# SPH20291（天然本体 = 全标记的 8 个标记原子全部被天然替换，tier-8）
NATURAL_FORMULA = "C₁₃₄H₁₉₈N₂₈O₃₅S₂"

SAMPLE_TYPES = {
    "sp": {"elements": LABELED_COLS, "formula": BODY_FORMULA_LABELED, "label": "SPH20291-Isotope1（全标记）"},
    "std": {"elements": LABELED_COLS, "formula": NATURAL_FORMULA, "label": "SPH20291（天然）"},
}


# ---------------------------------------------------------------------------
# centroid 谱图 → 合成 profile（用高斯峰替代每个 centroid 峰）
# ---------------------------------------------------------------------------
def centroid_to_profile(mz: np.ndarray, intensity: np.ndarray,
                        resolution: float, bin_da: float = PROFILE_BIN_DA,
                        extend_sigma: float = PROFILE_K_SIGMA) -> Tuple[np.ndarray, np.ndarray]:
    """把 centroid 谱图转换为合成 profile 谱图（Gaussian 峰替代）。"""
    if len(mz) == 0:
        return np.array([]), np.array([])
    lo = mz.min() - extend_sigma * mz.max() / resolution / GAUSS_K
    hi = mz.max() + extend_sigma * mz.max() / resolution / GAUSS_K
    grid = np.arange(lo, hi + bin_da * 0.5, bin_da)
    profile = np.zeros(len(grid), dtype=float)
    for m, a in zip(mz, intensity):
        sigma = m / resolution / GAUSS_K
        if sigma <= 0:
            continue
        profile += a * np.exp(-0.5 * ((grid - m) / sigma) ** 2)
    return grid, profile


# ---------------------------------------------------------------------------
# 主流程
# ---------------------------------------------------------------------------
def run_deconv(input_path: str, z: int, rt_min: float, rt_max: float,
               resolution: float = RESOLUTION,
               elements: List[Tuple[str, str, int]] = None,
               sample_type: str = "sp",
               out_dir: str = None) -> dict:
    """运行解卷积 → 返回统一格式结果 dict。"""
    if elements is None:
        elements = SAMPLE_TYPES[sample_type]["elements"]
    if out_dir is None:
        out_dir = HERE

    # 1) 获取谱图（LCD 或 CSV）
    print(f"[1/4] 读取谱图：{input_path}")
    ext = os.path.splitext(input_path)[1].lower()

    if ext == '.lcd':
        # 从 LCD 导出谱图 CSV
        mz_raw, intensity_raw = lcd2csv.export_rt_average(input_path, rt_min, rt_max, ms_level=1)
        print(f"      RT [{rt_min}-{rt_max}]s 平均 → {len(mz_raw)} 个 centroid 峰")
        if len(mz_raw) == 0:
            raise ValueError("LCD 导出无数据")
    elif ext == '.csv':
        # 直接读取 CSV
        data = np.genfromtxt(input_path, delimiter=',', skip_header=1)
        mz_raw, intensity_raw = data[:, 0], data[:, 1]
        print(f"      CSV 直接读取 → {len(mz_raw)} 个峰")
    else:
        raise ValueError(f"不支持的文件格式: {ext}")

    if len(mz_raw) < 3:
        raise ValueError(f"谱图有效点数不足: {len(mz_raw)}")

    # 2) centroid → synthetic profile（让 NNLS 有足够采样点）
    print(f"[2/4] centroid → synthetic profile（bin={PROFILE_BIN_DA} Da, R={resolution}）")
    grid, profile = centroid_to_profile(mz_raw, intensity_raw, resolution)
    print(f"      {len(grid)} 个网格点，范围 {grid.min():.2f}-{grid.max():.2f} Da")

    # 保存临时 CSV 供 deconv 读取
    tmp_csv = os.path.join(out_dir, f"_tmp_profile_{sample_type}_z{z}.csv")
    with open(tmp_csv, 'w') as f:
        f.write("m/z,intensity\n")
        for m, a in zip(grid, profile):
            f.write(f"{m:.6f},{a:.8e}\n")

    # 3) 解卷积
    print(f"[3/4] 解卷积（z={z}, 样品类型={sample_type}）...")
    cols = [imp_mod.ElementColumn(e, i, c) for e, i, c in elements]
    result = deconv.deconvolve(
        cols, tmp_csv,
        z=z, resolution=resolution,
        ppm_range=10.0, ppm_step=1.0,
        with_baseline=True, merge_degenerate=True,
    )
    print(f"      offset={result['mass_offset_ppm']:+.1f} ppm, RSS={result['residual_rss']:.3e}, "
          f"groups={result['n_groups']}, species={result['n_species']}")

    # 4) 转换为 tier 格式
    print(f"[4/4] 转换为 9 档体系...")
    tier_result = groups_to_tiers(result, sample_type=sample_type)

    # 清理临时文件
    try:
        os.remove(tmp_csv)
    except OSError:
        pass

    return {
        "z": z,
        "input": input_path,
        "sample_type": sample_type,
        "rt": (rt_min, rt_max),
        "mode": result["mode"],
        "resolution": resolution,
        "offset_ppm": result["mass_offset_ppm"],
        "rss": result["residual_rss"],
        "n_groups": result["n_groups"],
        "n_species": result["n_species"],
        "tiers": tier_result,
        "raw_groups": result["groups"],
        "raw_species": result["species"],
    }


def groups_to_tiers(result: dict, sample_type: str = "sp") -> List[dict]:
    """把 deconv 的 group 列表转换为 9 档（tier）格式。

    对于 sp（全标记本体）：tier t = t 个标记原子被天然替换，t=0 是本体
    对于 std（天然本体）：tier t = t 个天然原子被标记替换，t=0 是天然本体
    """
    tiers = {t: {"tier": t, "content_pct": 0.0, "members": [], "n_species": 0} for t in range(9)}
    for g in result["groups"]:
        member_rank = g["members"][0]
        species = result["species"][member_rank - 1]
        mod_total = sum(c for el, iso, c in zip(species["elements"], species["isos"], species["counts"])
                        if iso != "natural")
        # sp: tier = 8 - mod_total (labeled body = tier 0)
        # std: tier = mod_total (natural body = tier 0)
        t = mod_total if sample_type == "std" else (8 - mod_total)
        if 0 <= t <= 8:
            tiers[t]["content_pct"] += g["relative_content"]
            tiers[t]["n_species"] += g["n_members"]
            tiers[t]["members"].append(species["name"])
    return [tiers[t] for t in range(9)]


# ---------------------------------------------------------------------------
# xlsx 输出
# ---------------------------------------------------------------------------
def write_xlsx(results: List[dict], out_path: str):
    """把多个 result 写入 xlsx（格式与 run_all.py 一致，1 位小数）。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    WORK = Font(bold=True)
    THIN = Side(style="thin", color="BFBFBF")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    CENTER = Alignment(horizontal="center", vertical="center")
    LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

    def put(ws, r, c, v, bold=False, fill=None, align=CENTER, num=None):
        cell = ws.cell(row=r, column=c, value=v)
        if bold:
            cell.font = WORK
        if fill:
            cell.fill = fill
        cell.alignment = align
        if num:
            cell.number_format = num
        cell.border = BORDER
        return cell

    wb = Workbook()
    ws = wb.active
    ws.title = "解卷积_杂质含量"

    hdr = ["样品", "类型", "z", "RT窗口", "档", "含量%", "成员数", "质量偏移ppm", "RSS"]
    widths = [20, 14, 6, 18, 8, 10, 8, 14, 14]
    for i, (h, w) in enumerate(zip(hdr, widths), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = CENTER
        cell.border = BORDER

    r = 2
    for res in results:
        sample = os.path.basename(res["input"])
        type_label = SAMPLE_TYPES.get(res["sample_type"], {}).get("label", res["sample_type"])
        for tier in res["tiers"]:
            put(ws, r, 1, sample)
            put(ws, r, 2, type_label)
            put(ws, r, 3, f"z{res['z']}")
            put(ws, r, 4, f"{res['rt'][0]:.1f}-{res['rt'][1]:.1f}s")
            put(ws, r, 5, tier["tier"], bold=(tier["tier"] == 0))
            put(ws, r, 6, round(tier["content_pct"], 1), num="0.0",
                bold=(tier["tier"] == 0))
            put(ws, r, 7, len(tier["members"]))
            put(ws, r, 8, round(res["offset_ppm"], 1), num="0.0")
            put(ws, r, 9, f"{res['rss']:.2e}")
            if tier["tier"] == 0:
                for c in range(1, 10):
                    ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="E2EFDA")
            r += 1

    wb.save(out_path)
    print(f"\n写入 xlsx：{out_path}")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def parse_args(argv=None):
    p = argparse.ArgumentParser(
        prog="deconv_impurity",
        description="基于解卷积的同位素杂质含量计算（.lcd/.csv → xlsx）",
    )
    p.add_argument("--input", nargs="+", required=True,
                   help="输入文件路径（.lcd 或 .csv，可多个，空格分隔）")
    p.add_argument("--sample-type", nargs="+", default=["sp"],
                   choices=["sp", "std"],
                    help="样品类型（sp=全标记，std=天然；默认 sp）")
    p.add_argument("--z", nargs="+", type=int, default=[3],
                   help="电荷态（默认 3，可多个如 3 4）")
    p.add_argument("--rt", nargs=2, type=float, default=None,
                   help="RT 窗口 [MIN, MAX] 秒（仅 .lcd 有效；默认自动检测 TIC 峰值 ±5s）")
    p.add_argument("--resolution", type=float, default=RESOLUTION,
                   help=f"仪器分辨率（默认 {RESOLUTION}）")
    p.add_argument("--out-xlsx", default=None,
                   help="输出 xlsx 路径（默认：deconv_impurity_output.xlsx）")
    p.add_argument("--out-dir", default=None,
                   help="输出目录（默认：本目录）")
    return p.parse_args(argv)


def main(argv=None):
    args = parse_args(argv)
    out_dir = args.out_dir or HERE
    os.makedirs(out_dir, exist_ok=True)

    all_results = []
    for input_path in args.input:
        print(f"\n{'='*60}")
        print(f"处理：{input_path}")
        print(f"{'='*60}")
        if not os.path.exists(input_path):
            print(f"  跳过（文件不存在）: {input_path}")
            continue

        # 确定样品类型
        # 如果只有一个 sample type，用于所有文件；否则一一对应
        idx = args.input.index(input_path)
        if len(args.sample_type) == 1:
            sample_type = args.sample_type[0]
        else:
            sample_type = args.sample_type[idx] if idx < len(args.sample_type) else args.sample_type[-1]

        ext = os.path.splitext(input_path)[1].lower()

        if ext == '.csv':
            # CSV 文件无需 RT 窗口
            rt_min, rt_max = 0.0, 0.0
            print(f"  样品类型：{SAMPLE_TYPES[sample_type]['label']}")
        elif ext == '.lcd':
            # LCD 需要 RT 窗口
            if args.rt:
                rt_min, rt_max = args.rt
            else:
                rows = lcd2csv.list_scans(input_path)
                ms1 = [(i, rt) for i, rt, ms, n in rows if ms == 1]
                if not ms1:
                    print(f"  跳过（无 MS1 扫描）")
                    continue
                tics = []
                for i, rt in ms1[::5]:
                    m, it = lcd2csv.export_scan(input_path, i)
                    tics.append((rt, float(np.sum(it))))
                rt_vals, tic_vals = zip(*tics)
                tic_sm = np.convolve(tic_vals, np.ones(5) / 5, mode='same')
                peak_idx = np.argmax(tic_sm)
                peak_rt = rt_vals[peak_idx]
                rt_min, rt_max = peak_rt - 5.0, peak_rt + 5.0
                print(f"  自动 RT 窗口：{rt_min:.1f}-{rt_max:.1f} s（TIC 峰值 ~{peak_rt:.1f} s）")

        for z in args.z:
            print(f"\n--- z={z} ---")
            try:
                result = run_deconv(input_path, z, rt_min, rt_max,
                                    resolution=args.resolution,
                                    sample_type=sample_type, out_dir=out_dir)
                print(f"\n{'档':>3}  {'含量%':>8}  {'成员数':>6}")
                for t in result["tiers"]:
                    print(f"{t['tier']:>3}  {t['content_pct']:>8.1f}  {len(t['members']):>6}")
                all_results.append(result)
            except Exception as e:
                print(f"  错误：{e}")
                import traceback
                traceback.print_exc()

    if all_results:
        out_xlsx = args.out_xlsx or os.path.join(out_dir, "deconv_impurity_output.xlsx")
        write_xlsx(all_results, out_xlsx)

    print("\n完成。")


if __name__ == "__main__":
    sys.exit(main())
