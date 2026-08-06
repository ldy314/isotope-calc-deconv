# -*- coding: utf-8 -*-
"""
生成 同位素取代杂质计算器 (含 CSV 输出) .xlsm

与 ExactMass_Impurities.xlsm 相同的输入格式（三行元素表 + 级联下拉），
另加参数区：
  B9  电荷数 z（默认 1）
  B10 质量精度（峰合并容差 Da，默认 0.001）
  B11 丰度阈值 %（默认 0.05）
按钮"列举并计算" → 枚举杂质 + 逐杂质调用 theo.py 算同位素峰 → 输出 CSV
"""
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = r'D:\code test\chem\同位素计算及解卷积\ExactMass_Impurities_Calc.xlsx'
ISO_JSON = r'D:\code test\chem\同位素计算及解卷积\isotope_data.json'

wb = openpyxl.Workbook()

# ---------------------------------------------------------------------------
# 样式常量
# ---------------------------------------------------------------------------
FONT_TITLE = Font(name='微软雅黑', size=14, bold=True, color='1F4E79')
FONT_LABEL = Font(name='微软雅黑', size=11, bold=True, color='1F4E79')
FONT_BODY = Font(name='微软雅黑', size=11)
FONT_HEADER = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
FILL_HEADER = PatternFill('solid', fgColor='2E75B6')
FILL_INPUT = PatternFill('solid', fgColor='FFF2CC')
FILL_OUTPUT = PatternFill('solid', fgColor='E2EFDA')
FILL_ISO = PatternFill('solid', fgColor='DDEBF7')
THIN = Side(style='thin', color='B4B4B4')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal='center', vertical='center')

# ---------------------------------------------------------------------------
# IsotopeData sheet
# ---------------------------------------------------------------------------
ws_iso = wb.active
ws_iso.title = 'IsotopeData'
with open(ISO_JSON, encoding='utf-8') as f:
    iso_rows = json.load(f)

headers = ['元素', '质量数', '精确质量 (u)', '天然丰度']
for c, h in enumerate(headers, 1):
    cell = ws_iso.cell(row=1, column=c, value=h)
    cell.font = FONT_HEADER
    cell.fill = FILL_HEADER
    cell.alignment = CENTER
    cell.border = BORDER

for r, row in enumerate(iso_rows, 2):
    ws_iso.cell(row=r, column=1, value=row['element']).font = FONT_BODY
    ws_iso.cell(row=r, column=2, value=row['mass_number']).font = FONT_BODY
    c3 = ws_iso.cell(row=r, column=3, value=row['mass'])
    c3.font = FONT_BODY
    c3.number_format = '0.000000000'
    c4 = ws_iso.cell(row=r, column=4, value=row['abundance'])
    c4.font = FONT_BODY
    c4.number_format = '0.000000'
    for c in range(1, 5):
        ws_iso.cell(row=r, column=c).border = BORDER
        ws_iso.cell(row=r, column=c).fill = FILL_ISO
        ws_iso.cell(row=r, column=c).alignment = CENTER

for col, width in zip('ABCD', [10, 10, 18, 12]):
    ws_iso.column_dimensions[col].width = width

ELEM_LIST = []
for row in iso_rows:
    if row['element'] not in ELEM_LIST:
        ELEM_LIST.append(row['element'])

# 同位素下拉辅助区
ISO_HELPER_START_COL = 6
for idx, elem in enumerate(ELEM_LIST):
    hcol = ISO_HELPER_START_COL + idx
    HL = get_column_letter(hcol)
    head = ws_iso.cell(row=1, column=hcol, value=elem)
    head.font = FONT_HEADER
    head.fill = FILL_HEADER
    head.alignment = CENTER
    head.border = BORDER
    opts = ['natural'] + [str(r['mass_number']) for r in iso_rows if r['element'] == elem]
    for j, opt in enumerate(opts, start=2):
        ccell = ws_iso.cell(row=j, column=hcol, value=opt)
        ccell.font = FONT_BODY
        ccell.alignment = CENTER
        ccell.border = BORDER
    last_row = 1 + len(opts)
    wb.defined_names.add(openpyxl.workbook.defined_name.DefinedName(
        f'iso_{elem}',
        attr_text=f"IsotopeData!${HL}$2:${HL}${last_row}",
    ))
ws_iso.column_dimensions['F'].width = 10
ws_iso.column_dimensions['G'].width = 10

# ---------------------------------------------------------------------------
# Calculator sheet
# ---------------------------------------------------------------------------
ws = wb.create_sheet('Calculator')

ws.merge_cells('A1:J1')
c = ws['A1']
c.value = '同位素取代杂质计算器（枚举 + 逐杂质同位素峰 → CSV）'
c.font = FONT_TITLE
c.alignment = Alignment(horizontal='left', vertical='center')

# --- 输入区（三行元素表） ---
ws['A3'] = '元素'
ws['A3'].font = FONT_LABEL
ws['A4'] = '同位素种类'
ws['A4'].font = FONT_LABEL
ws['A5'] = '原子个数'
ws['A5'].font = FONT_LABEL

# 默认示例：CH2D3
defaults = [
    ('C', 'natural', 1),
    ('H', 'natural', 2),
    ('H', '2', 3),
]
N_COLS = 26

for i in range(N_COLS):
    col = i + 2
    L = get_column_letter(col)
    cell3 = ws[f'{L}3']
    if i < len(defaults):
        cell3.value = defaults[i][0]
    cell3.fill = FILL_INPUT
    cell3.font = FONT_BODY
    cell3.alignment = CENTER
    cell3.border = BORDER
    dv3 = openpyxl.worksheet.datavalidation.DataValidation(
        type='list', formula1=f'="{",".join(ELEM_LIST)}"', allow_blank=True,
        showDropDown=False, showErrorMessage=True)
    dv3.error = '请从下拉列表选择元素'
    dv3.errorTitle = '无效元素'
    ws.add_data_validation(dv3)
    dv3.add(cell3)

    cell4 = ws[f'{L}4']
    if i < len(defaults):
        cell4.value = defaults[i][1]
    cell4.fill = FILL_INPUT
    cell4.font = FONT_BODY
    cell4.alignment = CENTER
    cell4.border = BORDER
    dv4 = openpyxl.worksheet.datavalidation.DataValidation(
        type='list',
        formula1=f'=IF({L}3="","",INDIRECT("iso_"&{L}3))',
        allow_blank=True,
        showDropDown=False, showErrorMessage=True)
    dv4.error = 'natural 或质量数（如 2=氘）'
    dv4.errorTitle = '无效同位素'
    ws.add_data_validation(dv4)
    dv4.add(cell4)

    cell5 = ws[f'{L}5']
    if i < len(defaults):
        cell5.value = defaults[i][2]
    cell5.fill = FILL_INPUT
    cell5.font = FONT_BODY
    cell5.alignment = CENTER
    cell5.border = BORDER
    cell5.number_format = '0'

for col in range(2, 2 + N_COLS):
    ws.column_dimensions[get_column_letter(col)].width = 9

# --- 分子式显示（B7 自动，含电荷） ---
SUB_MAP_D = {'0': '₀', '1': '₁', '2': '₂', '3': '₃', '4': '₄',
             '5': '₅', '6': '₆', '7': '₇', '8': '₈', '9': '₉'}
SUB_MAP_U = {'0': '⁰', '1': '¹', '2': '²', '3': '³', '4': '⁴',
             '5': '⁵', '6': '⁶', '7': '⁷', '8': '⁸', '9': '⁹'}


def _sub_formula(ref: str, mapping: dict) -> str:
    f = ref
    for d, rep in mapping.items():
        f = f'SUBSTITUTE({f},"{d}","{rep}")'
    return f


for i in range(N_COLS):
    col = i + 2
    L = get_column_letter(col)
    cnt_fmt = _sub_formula(f'TEXT({L}5,"0")', SUB_MAP_D)
    iso_fmt = _sub_formula(f'TEXT({L}4,"0")', SUB_MAP_U)
    cell6 = ws[f'{L}6']
    cell6.value = (
        f'=IF({L}5="","",'
        f'IF(OR({L}4="",{L}4="natural"),'
        f'{L}3&{cnt_fmt},'
        f'{iso_fmt}&{L}3&{cnt_fmt}))'
    )
    cell6.font = Font(name='微软雅黑', size=8, color='999999')

ws['A7'] = '分子式（自动）'
ws['A7'].font = FONT_LABEL
ws.merge_cells('B7:J7')
ws['B7'].font = FONT_BODY

CHARGE_SUP = {
    '0': '⁰', '1': '¹', '2': '²', '3': '³', '4': '⁴',
    '5': '⁵', '6': '⁶', '7': '⁷', '8': '⁸', '9': '⁹',
}
abs_z = _sub_formula('TEXT(ABS(B9),"0")', CHARGE_SUP)
charge_abs = (
    f'IF(ABS(B9)=1,'
    f'IF(B9>0,"⁺","⁻"),'
    f'{abs_z}&IF(B9>0,"⁺","⁻"))'
)
charge_suffix = f'IF(OR(B9="",B9=0),"",{charge_abs})'
ws['B7'].value = (
    '=CONCATENATE(B6, C6, D6, E6, F6, G6, H6, I6, J6, K6, L6, M6, N6, O6, P6, Q6, R6, S6, T6, U6, V6, W6, X6, Y6, Z6, AA6)'
    ' & ' + charge_suffix
)

# --- 参数区：z / tol / min-ab（与 ExactMass_Calculator_v2 一致） ---
ws['A9'] = '电荷数 z'
ws['A9'].font = FONT_LABEL
ws['B9'] = 1
ws['B9'].font = FONT_BODY
ws['B9'].alignment = CENTER
ws['B9'].border = BORDER
ws['B9'].fill = FILL_INPUT

ws['A10'] = '质量精度 (Da)'
ws['A10'].font = FONT_LABEL
ws['B10'] = 0.001
ws['B10'].font = FONT_BODY
ws['B10'].alignment = CENTER
ws['B10'].border = BORDER
ws['B10'].fill = FILL_INPUT
ws['B10'].number_format = '0.000'

ws['A11'] = '丰度阈值'
ws['A11'].font = FONT_LABEL
ws['B11'] = 0.05
ws['B11'].font = FONT_BODY
ws['B11'].alignment = CENTER
ws['B11'].border = BORDER
ws['B11'].fill = FILL_INPUT
ws['B11'].number_format = '0.00"%"'

# --- 使用说明 ---
ws['A13'] = '使用说明：'
ws['A13'].font = FONT_LABEL
notes = [
    '1. 输入同位素修饰分子式：第3行元素、第4行同位素种类（natural 或质量数如 2=氘）、第5行原子个数',
    '2. 参数区：第9行电荷数 z，第10行质量精度（峰合并容差 Da），第11行丰度阈值 %（低于不显示）',
    '3. 点击"列举并计算"按钮：枚举所有同位素取代杂质（含全天然，不含输入本身），',
    '   并对每个杂质用 ExactMass_Calculator_v2 的引擎计算其理论同位素峰（前20峰）',
    '4. 结果：① Excel 输出区显示杂质列表（分子式/元素/同位素种类/原子个数）',
    '   ② 同目录生成 impurity_peaks_<时间戳>.csv（每行=杂质-峰对，UTF-8 可直接用 Excel 打开）',
    '5. natural 列数字只增不减（底线=输入值）；同元素同同位素合并；多修饰列按组合枚举',
    '6. 需要本机安装 Python 与 molmass 库：pip install molmass',
]
for i, note in enumerate(notes):
    cell = ws[f'A{14 + i}']
    cell.value = note
    cell.font = Font(name='微软雅黑', size=10, color='595959')

# --- 输出区 ---
ws['A23'] = '输出：同位素取代杂质（每杂质 4 行：分子式 / 元素 / 同位素种类 / 原子个数）'
ws['A23'].font = FONT_LABEL

OUT_ROWS = 600
for r in range(24, 24 + OUT_ROWS):
    ws.cell(row=r, column=1).font = Font(name='微软雅黑', size=10, bold=True, color='1F4E79')
    for c in range(2, 12):
        cell = ws.cell(row=r, column=c)
        cell.font = FONT_BODY
        cell.border = BORDER
        cell.fill = FILL_OUTPUT
        cell.alignment = CENTER
    ws.cell(row=r, column=2).alignment = Alignment(horizontal='left', vertical='center')
    ws.cell(row=r, column=3).alignment = Alignment(horizontal='left', vertical='center')

ws.column_dimensions['A'].width = 20
ws.column_dimensions['B'].width = 12
ws.column_dimensions['C'].width = 12
ws.column_dimensions['D'].width = 12
ws.column_dimensions['E'].width = 12
ws.column_dimensions['F'].width = 12

# --- 按钮说明 ---
ws['A633'] = '▶ 按钮"列举并计算"由构建脚本自动放置（E40 区域）'
ws['A633'].font = Font(name='微软雅黑', size=10, bold=True, color='C00000')

wb.save(OUT)
print(f'已保存: {OUT}（此为 .xlsx 中间格式，随后由注入脚本注入宏并另存为 .xlsm）')
